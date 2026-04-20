#!/usr/bin/env python3
"""Generate comprehensive Vietnamese project guide as Word document.
Run: cd semisub-cxr && .venv/bin/python scripts/gen_guide.py
Output: outputs/huong_dan_du_an_chi_tiet.docx
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def main():
    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    # ============================================================
    # TITLE
    # ============================================================
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "HUONG DAN CHI TIET DU AN\n"
        "Semi-Supervised Multi-Label Classification\n"
        "of Chest X-Ray Images"
    )
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Tai lieu giai thich toan bo ky thuat, workflow va code tu A-Z")
    r.font.size = Pt(14)
    r.italic = True

    doc.add_page_break()

    build_toc(doc)
    doc.add_page_break()
    build_section1(doc)
    build_section2(doc)
    build_section3(doc)
    build_section4(doc)
    build_section5(doc)
    build_section6(doc)
    build_section7(doc)
    build_section8(doc)
    build_section9(doc)
    build_section10(doc)

    # Save
    out_dir = "outputs"
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, "huong_dan_du_an_chi_tiet.docx")
    doc.save(path)
    print(f"Document saved to: {path}")


def build_toc(doc):
    doc.add_heading("MUC LUC", level=1)
    items = [
        "1. Tong quan du an (Project Overview)",
        "2. Dataset CheXpert - Du lieu va Tien xu ly",
        "3. Kien truc mo hinh (Model Architecture)",
        "4. Quy trinh chia du lieu (Split Generation)",
        "5. Training Pipeline - Quy trinh huan luyen",
        "6. Pseudo-Labeling - Gan nhan gia",
        "7. MC Dropout - Uoc luong do bat dinh",
        "8. Evaluation - Danh gia mo hinh",
        "9. Cau truc code chi tiet",
        "10. Cau hoi thuyet trinh thuong gap",
    ]
    for item in items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)



def build_section1(doc):
    doc.add_heading("1. TONG QUAN DU AN (Project Overview)", level=1)

    doc.add_paragraph(
        "Du an nay nghien cuu bai toan phan loai da nhan (multi-label classification) "
        "tren anh X-quang nguc (Chest X-Ray) su dung phuong phap hoc ban giam sat "
        "(semi-supervised learning). Van de chinh: trong y te, viec gan nhan du lieu "
        "can bac si chuyen mon, rat ton kem va mat thoi gian. Du an nay chi su dung "
        "5% du lieu co nhan de huan luyen, roi tan dung 95% du lieu khong nhan "
        "thong qua ky thuat pseudo-labeling."
    )

    doc.add_heading("1.1. Bai toan", level=2)
    doc.add_paragraph(
        "Input: Anh X-quang nguc (224x224 pixels, 3 kenh mau RGB).\n"
        "Output: 6 xac suat tuong ung voi 6 benh ly:\n"
        "  - Cardiomegaly (tim to)\n"
        "  - Pleural Effusion (tran dich mang phoi)\n"
        "  - Pneumothorax (tran khi mang phoi)\n"
        "  - Consolidation (dong dac phoi)\n"
        "  - Atelectasis (xep phoi)\n"
        "  - Edema (phu phoi)\n\n"
        "Day la bai toan MULTI-LABEL: mot anh co the co nhieu benh cung luc, "
        "khac voi multi-class (chi 1 nhan duy nhat)."
    )

    doc.add_heading("1.2. 4 Setting thi nghiem", level=2)
    doc.add_paragraph(
        "Setting 1 - Supervised Baseline:\n"
        "  Train ResNet-50 (ImageNet pretrained) chi tren 5% du lieu co nhan.\n"
        "  Day la baseline de so sanh.\n\n"
        "Setting 2 - LVM-Med Fine-tuning:\n"
        "  Thay backbone ImageNet bang LVM-Med (pretrained tren anh y te).\n"
        "  Ky vong: feature tot hon vi domain-specific.\n\n"
        "Setting 3 - Pseudo-Labeling:\n"
        "  Dung mo hinh da train (teacher) de gan nhan gia cho du lieu khong nhan.\n"
        "  Train mo hinh moi (student) tren ca du lieu that + nhan gia.\n\n"
        "Setting 4 - Uncertainty-Filtered Pseudo-Labeling:\n"
        "  Giong Setting 3 nhung them MC Dropout de loc bo nhan gia khong dang tin cay.\n"
        "  Chi giu lai nhan gia co do tin cay cao VA do bat dinh thap."
    )

    doc.add_heading("1.3. Workflow tong the", level=2)
    doc.add_paragraph(
        "Buoc 1: Load CheXpert dataset (223K anh) -> loc frontal views (191K anh)\n"
        "Buoc 2: Chia 5% labeled (9,551) + 95% unlabeled (181,476) bang stratified split\n"
        "Buoc 3: Train teacher model tren 5% labeled -> best checkpoint\n"
        "Buoc 4: Teacher chay inference tren 181K unlabeled -> xac suat du doan\n"
        "Buoc 5: Ap dung threshold (0.7) -> tao pseudo-labels\n"
        "Buoc 6: Train student model tren labeled + pseudo-labeled\n"
        "Buoc 7: (Optional) MC Dropout -> loc uncertainty -> train student tot hon\n"
        "Buoc 8: Evaluate tat ca models, so sanh AUROC"
    )



def build_section2(doc):
    doc.add_heading("2. DATASET CHEXPERT - DU LIEU VA TIEN XU LY", level=1)

    doc.add_heading("2.1. CheXpert la gi?", level=2)
    doc.add_paragraph(
        "CheXpert (Irvin et al., 2019) la bo du lieu X-quang nguc lon tu Benh vien Stanford:\n"
        "  - 224,316 anh X-quang tu 65,240 benh nhan\n"
        "  - 14 nhan benh ly duoc trich xuat tu bao cao X-quang bang NLP\n"
        "  - Moi nhan co 4 gia tri: positive (1), negative (0), uncertain (-1), blank (NaN)\n\n"
        "Du an chi dung 6/14 nhan benh ly quan trong nhat (xem Section 1.1)."
    )

    doc.add_heading("2.2. Xu ly nhan Uncertain (-1)", level=2)
    doc.add_paragraph(
        "Van de: Nhan -1 nghia la 'khong chac chan' - bac si khong ro benh nhan co benh hay khong.\n\n"
        "Giai phap - U-Zeros policy: Map tat ca -1 thanh 0 (coi nhu khong co benh).\n"
        "Ly do: Day la cach tiep can an toan (conservative), tranh false positive.\n\n"
        "Code thuc hien trong CheXpertDataset.__init__():\n"
        "  fill_value = 0.0  # uncertain_policy == 'zeros'\n"
        "  df[col] = df[col].fillna(0.0)  # NaN -> 0\n"
        "  df[col] = df[col].replace(-1.0, fill_value)  # -1 -> 0"
    )

    doc.add_heading("2.3. Loc View (Frontal Only)", level=2)
    doc.add_paragraph(
        "CheXpert co 2 loai view:\n"
        "  - Frontal (AP/PA): Chup tu phia truoc, la view chuan de chan doan\n"
        "  - Lateral: Chup tu ben hong\n\n"
        "Du an chi dung Frontal views vi:\n"
        "  1. Day la view pho bien nhat trong lam sang\n"
        "  2. Giam nhieu du lieu, tap trung vao view co gia tri chan doan cao\n"
        "  3. Sau khi loc: 191,027 frontal views tu 223,414 tong"
    )

    doc.add_heading("2.4. Image Transforms", level=2)
    doc.add_paragraph(
        "TRAINING transforms (data augmentation):\n"
        "  1. Resize(224, 224) - resize ve kich thuoc chuan cua ResNet\n"
        "  2. RandomHorizontalFlip() - lat ngang ngau nhien (tang da dang)\n"
        "  3. RandomRotation(10) - xoay ngau nhien +-10 do\n"
        "  4. ToTensor() - chuyen PIL Image thanh PyTorch tensor [0,1]\n"
        "  5. Normalize(mean, std) - chuan hoa theo ImageNet statistics\n\n"
        "EVALUATION transforms (khong augmentation):\n"
        "  1. Resize(224, 224)\n"
        "  2. ToTensor()\n"
        "  3. Normalize(mean, std)\n\n"
        "Normalization values (ImageNet standard):\n"
        "  mean = [0.485, 0.456, 0.406]\n"
        "  std  = [0.229, 0.224, 0.225]\n\n"
        "Tai sao dung ImageNet normalization cho anh y te?\n"
        "Vi backbone ResNet-50 duoc pretrain tren ImageNet, nen input phai "
        "duoc normalize giong nhu khi pretrain de feature extraction hoat dong tot."
    )



def build_section3(doc):
    doc.add_heading("3. KIEN TRUC MO HINH (Model Architecture)", level=1)

    doc.add_heading("3.1. Tong quan kien truc", level=2)
    doc.add_paragraph(
        "Mo hinh gom 3 thanh phan noi tiep:\n\n"
        "  Input Image (224x224x3)\n"
        "       |\n"
        "  [ResNet-50 Backbone] -> Feature vector (2048-dim)\n"
        "       |\n"
        "  [Dropout(p=0.5)] -> Regularization\n"
        "       |\n"
        "  [Linear(2048, 6)] -> 6 raw logits\n"
        "       |\n"
        "  sigmoid() -> 6 probabilities [0, 1]\n\n"
        "QUAN TRONG: Mo hinh tra ve RAW LOGITS (chua qua sigmoid).\n"
        "Sigmoid duoc ap dung BEN NGOAI trong loss function va evaluation."
    )

    doc.add_heading("3.2. ResNet-50 Backbone", level=2)
    doc.add_paragraph(
        "ResNet-50 la gi?\n"
        "  - Convolutional Neural Network voi 50 layers\n"
        "  - Su dung Residual Connections (skip connections) de giai quyet vanishing gradient\n"
        "  - Output: feature vector 2048 chieu\n\n"
        "Transfer Learning:\n"
        "  - Backbone duoc pretrain tren ImageNet (1.2 trieu anh, 1000 classes)\n"
        "  - Lop cuoi (fc layer) bi LOAI BO, thay bang nn.Identity()\n"
        "  - Chi giu lai phan feature extraction\n"
        "  - Fine-tune toan bo backbone tren du lieu CheXpert\n\n"
        "Code trong backbone_factory.py:\n"
        "  model = models.resnet50(weights=ResNet50_Weights.IMAGENET1K_V1)\n"
        "  model.fc = nn.Identity()  # Loai bo classification head\n\n"
        "Registry Pattern:\n"
        "  Du an dung registry pattern de de dang them backbone moi:\n"
        "  register_backbone('resnet50_imagenet', _load_resnet50_imagenet)\n"
        "  register_backbone('resnet50_lvmmed', _load_resnet50_lvmmed)\n"
        "  -> Goi create_backbone('resnet50_imagenet') se tra ve backbone tuong ung"
    )

    doc.add_heading("3.3. Dropout Layer", level=2)
    doc.add_paragraph(
        "Dropout(p=0.5) nghia la:\n"
        "  - Trong luc TRAINING: ngau nhien tat 50% neurons -> chong overfitting\n"
        "  - Trong luc INFERENCE (eval): tat ca neurons hoat dong binh thuong\n"
        "  - NGOAI TRU khi dung MC Dropout: bat dropout trong inference de uoc luong uncertainty\n\n"
        "Tai sao p=0.5?\n"
        "  Day la gia tri mac dinh pho bien, can bang giua regularization va capacity."
    )

    doc.add_heading("3.4. Classification Head", level=2)
    doc.add_paragraph(
        "nn.Linear(2048, 6):\n"
        "  - Input: 2048 features tu backbone\n"
        "  - Output: 6 logits (1 cho moi benh ly)\n"
        "  - KHONG co sigmoid o day (ap dung ben ngoai)\n\n"
        "Tai sao tra ve logits thay vi probabilities?\n"
        "  Vi BCEWithLogitsLoss ket hop sigmoid + BCE trong 1 buoc,\n"
        "  numerically stable hon viec tinh sigmoid roi BCE rieng."
    )



def build_section4(doc):
    doc.add_heading("4. QUY TRINH CHIA DU LIEU (Split Generation)", level=1)

    doc.add_heading("4.1. Van de", level=2)
    doc.add_paragraph(
        "Can chia 191,027 frontal training samples thanh:\n"
        "  - Labeled subset (5%): 9,551 samples co nhan -> dung de train\n"
        "  - Unlabeled subset (95%): 181,476 samples -> dung cho pseudo-labeling\n\n"
        "Yeu cau:\n"
        "  1. Moi class phai co it nhat 1 positive sample trong labeled subset\n"
        "  2. Phan chia phai REPRODUCIBLE (cung seed -> cung ket qua)\n"
        "  3. Phan chia phai STRATIFIED (giu ty le class balance)"
    )

    doc.add_heading("4.2. MultilabelStratifiedShuffleSplit", level=2)
    doc.add_paragraph(
        "Thu vien: iterative-stratification (iterstrat)\n\n"
        "Tai sao khong dung random split binh thuong?\n"
        "  Vi du lieu CheXpert rat IMBALANCED:\n"
        "  - Pleural Effusion: rat nhieu positive\n"
        "  - Pneumothorax: rat it positive\n"
        "  Neu random split, co the labeled subset khong co positive nao cho Pneumothorax!\n\n"
        "MultilabelStratifiedShuffleSplit dam bao:\n"
        "  - Ty le positive/negative cua MOI CLASS duoc giu nguyen\n"
        "  - Trong ca labeled va unlabeled subset\n\n"
        "Retry logic:\n"
        "  Neu split khong dat min_positive_per_class, thu lai voi seed khac (toi da 10 lan)."
    )

    doc.add_heading("4.3. Persistence", level=2)
    doc.add_paragraph(
        "Split duoc luu thanh JSON file:\n"
        "  outputs/splits/split_r0.05_s42.json\n\n"
        "Chua:\n"
        "  - labeled_indices: [0, 5, 12, ...] (9,551 indices)\n"
        "  - unlabeled_indices: [1, 2, 3, ...] (181,476 indices)\n"
        "  - per_class_positive_counts: so positive moi class\n"
        "  - generation_timestamp\n\n"
        "Lan chay sau se LOAD split cu thay vi tao moi -> dam bao reproducibility."
    )


def build_section5(doc):
    doc.add_heading("5. TRAINING PIPELINE - QUY TRINH HUAN LUYEN", level=1)

    doc.add_heading("5.1. Loss Function: Masked BCE with Logits", level=2)
    doc.add_paragraph(
        "Cong thuc Binary Cross-Entropy:\n"
        "  BCE(z, y) = -[y * log(sigmoid(z)) + (1-y) * log(1-sigmoid(z))]\n\n"
        "Trong do:\n"
        "  z = raw logit tu model\n"
        "  y = ground truth label (0 hoac 1)\n"
        "  sigmoid(z) = 1 / (1 + exp(-z))\n\n"
        "MASKED BCE:\n"
        "  L = sum(mask * BCE) / sum(mask)\n\n"
        "  mask[i,c] = 1.0: class c cua sample i THAM GIA vao loss\n"
        "  mask[i,c] = 0.0: class c cua sample i BI BO QUA\n\n"
        "Khi nao dung mask?\n"
        "  - Labeled samples: mask = all 1s (tat ca class deu tham gia)\n"
        "  - Pseudo-labeled samples: mask = 1 chi cho class duoc accept,\n"
        "    0 cho class bi reject -> tranh hoc tu nhan gia sai"
    )

    doc.add_heading("5.2. Optimizer va Scheduler", level=2)
    doc.add_paragraph(
        "Optimizer: Adam\n"
        "  - learning_rate = 0.0001 (1e-4)\n"
        "  - weight_decay = 0.0001 (L2 regularization)\n"
        "  - Adam tu dong dieu chinh learning rate cho tung parameter\n\n"
        "Scheduler: Cosine Annealing\n"
        "  - Giam learning rate theo hinh cosine tu lr_max xuong 0\n"
        "  - T_max = num_epochs (30)\n"
        "  - Giup model hoi tu tot hon o cuoi training\n\n"
        "Tai sao Adam + Cosine?\n"
        "  - Adam: hoi tu nhanh, it nhay cam voi learning rate\n"
        "  - Cosine: giam lr tu tu, tranh overfitting o cuoi"
    )

    doc.add_heading("5.3. Training Loop chi tiet", level=2)
    doc.add_paragraph(
        "Moi epoch gom:\n\n"
        "1. TRAIN PHASE:\n"
        "   for batch in train_loader:\n"
        "     images, labels, mask = unpack_batch(batch)\n"
        "     logits = model(images)           # Forward pass\n"
        "     loss = masked_bce(logits, labels, mask)  # Tinh loss\n"
        "     optimizer.zero_grad()            # Xoa gradient cu\n"
        "     loss.backward()                  # Backward pass (tinh gradient)\n"
        "     optimizer.step()                 # Cap nhat weights\n\n"
        "2. VALIDATION PHASE:\n"
        "   model.eval()  # Tat dropout, batch norm dung running stats\n"
        "   for batch in val_loader:\n"
        "     logits = model(images)\n"
        "     probs = sigmoid(logits)\n"
        "     -> Thu thap tat ca predictions\n"
        "   -> Tinh AUROC tren toan bo validation set\n\n"
        "3. CHECKPOINT:\n"
        "   - Luon luu last_checkpoint.pt\n"
        "   - Neu macro_auroc >= best -> luu best_checkpoint.pt\n"
        "   - Checkpoint chua: model weights, optimizer state, epoch, config, metrics"
    )

    doc.add_heading("5.4. Overfitting va Early Stopping", level=2)
    doc.add_paragraph(
        "Ket qua thuc te: Best epoch = 6 (trong 30 epochs)\n"
        "  -> Mo hinh bat dau overfit sau epoch 6\n"
        "  -> best_checkpoint.pt luu tai epoch 6 voi macro_auroc = 0.8195\n\n"
        "Tai sao overfit?\n"
        "  - Chi co 9,551 labeled samples (5%) -> du lieu qua it\n"
        "  - ResNet-50 co 25 trieu parameters -> model qua lon\n"
        "  -> Day chinh la ly do can pseudo-labeling: tang du lieu training!"
    )



def build_section6(doc):
    doc.add_heading("6. PSEUDO-LABELING - GAN NHAN GIA", level=1)

    doc.add_heading("6.1. Y tuong chinh", level=2)
    doc.add_paragraph(
        "Van de: Chi co 5% du lieu co nhan, 95% khong co nhan.\n"
        "Y tuong: Dung mo hinh da train (teacher) de 'doan' nhan cho 95% con lai,\n"
        "roi train mo hinh moi (student) tren ca du lieu that + du lieu doan.\n\n"
        "TEACHER = mo hinh supervised baseline (da train xong, epoch 6, AUROC 0.82)\n"
        "STUDENT = mo hinh moi, khoi tao tu dau (from_scratch)\n\n"
        "Tai sao student khong copy weights tu teacher?\n"
        "  Vi student_init = 'from_scratch' -> train tu ImageNet pretrained weights\n"
        "  Muc dich: student hoc tu du lieu (labeled + pseudo) thay vi copy teacher"
    )

    doc.add_heading("6.2. Teacher Inference", level=2)
    doc.add_paragraph(
        "Buoc 1: Load teacher model tu best_checkpoint.pt\n"
        "Buoc 2: Chay inference tren 181,476 unlabeled samples\n"
        "Buoc 3: Thu thap sigmoid probabilities cho moi sample, moi class\n\n"
        "Code: run_teacher_inference(model, dataloader, device)\n"
        "  model.eval()  # Tat dropout\n"
        "  for batch in dataloader:\n"
        "    logits = model(images)\n"
        "    probs = sigmoid(logits)  # -> (batch_size, 6)\n"
        "  -> Ket qua: probabilities array shape (181476, 6)"
    )

    doc.add_heading("6.3. Threshold Strategy: PositiveOnlyStrategy", level=2)
    doc.add_paragraph(
        "Voi moi sample i va class c:\n"
        "  NEU prob[i,c] > threshold_c (0.7):\n"
        "    pseudo_label[i,c] = 1.0  (ACCEPT - teacher tin la positive)\n"
        "  NGUOC LAI:\n"
        "    pseudo_label[i,c] = NaN   (REJECT - khong du tin cay)\n\n"
        "TAI SAO chi gan nhan POSITIVE (1.0), khong gan NEGATIVE (0.0)?\n"
        "  Vi voi benh hiem (Pneumothorax), teacher co the du doan prob thap\n"
        "  cho ca benh nhan THUC SU co benh. Neu gan negative, se lam mat\n"
        "  nhung positive samples quy gia -> giam hieu qua.\n\n"
        "Ket qua thuc te voi threshold 0.7:\n"
        "  Cardiomegaly:      13,175 accepted (7.3%)\n"
        "  Pleural Effusion:  41,983 accepted (23.1%)\n"
        "  Pneumothorax:       3,884 accepted (2.1%)\n"
        "  Consolidation:         21 accepted (0.01%) <- Rat it!\n"
        "  Atelectasis:            56 accepted (0.03%) <- Rat it!\n"
        "  Edema:             14,221 accepted (7.8%)\n\n"
        "Nhan xet: Consolidation va Atelectasis gan nhu khong co pseudo-label\n"
        "  -> Teacher khong tu tin ve 2 class nay\n"
        "  -> Co the can ha threshold cho 2 class nay"
    )

    doc.add_heading("6.4. CombinedDataset", level=2)
    doc.add_paragraph(
        "Sau khi co pseudo-labels, tao CombinedDataset gop:\n"
        "  - 9,551 labeled samples (nhan that, mask = all 1s)\n"
        "  - 181,476 pseudo-labeled samples (nhan gia, mask theo accepted/rejected)\n\n"
        "Moi sample tra ve 4 gia tri:\n"
        "  (image, label_vector, loss_mask, is_pseudo)\n\n"
        "  Labeled sample:  mask = [1,1,1,1,1,1] -> tat ca class tham gia loss\n"
        "  Pseudo sample:   mask = [1,0,0,1,0,1] -> chi class accepted tham gia\n\n"
        "QUAN TRONG: Nhan that KHONG BAO GIO bi ghi de boi pseudo-label.\n"
        "  Labeled samples luon giu nguyen ground-truth labels."
    )

    doc.add_heading("6.5. Student Training", level=2)
    doc.add_paragraph(
        "Student model duoc train giong nhu supervised baseline,\n"
        "nhung tren CombinedDataset thay vi chi labeled subset.\n\n"
        "Trainer._unpack_batch() xu ly 2 format:\n"
        "  - 3-tuple (images, labels, indices) -> standard dataset, mask = None\n"
        "  - 4-tuple (images, labels, mask, is_pseudo) -> CombinedDataset\n\n"
        "Ky vong: Student se tot hon teacher vi co nhieu du lieu hon."
    )



def build_section7(doc):
    doc.add_heading("7. MC DROPOUT - UOC LUONG DO BAT DINH", level=1)

    doc.add_heading("7.1. Van de voi Pseudo-Labeling thong thuong", level=2)
    doc.add_paragraph(
        "Pseudo-labeling chi dua vao confidence (xac suat cao = tin cay).\n"
        "Nhung confidence CAO khong phai luc nao cung DUNG:\n"
        "  - Mo hinh co the 'overconfident' - du doan xac suat cao nhung sai\n"
        "  - Can them 1 thuoc do khac de danh gia do tin cay\n\n"
        "Giai phap: MC Dropout (Monte Carlo Dropout)\n"
        "  - Dua tren ly thuyet Bayesian Approximation (Gal & Ghahramani, 2016)\n"
        "  - Y tuong: Neu mo hinh THAT SU chac chan, thi du co bat dropout\n"
        "    nhieu lan, ket qua van NHAT QUAN (it thay doi)"
    )

    doc.add_heading("7.2. Cach hoat dong", level=2)
    doc.add_paragraph(
        "Buoc 1: BAT dropout trong inference (binh thuong dropout bi tat)\n"
        "  Code: _enable_dropout(model)\n"
        "    model.eval()  # Tat batch norm\n"
        "    for module in model.modules():\n"
        "      if isinstance(module, nn.Dropout):\n"
        "        module.train()  # Chi bat lai dropout\n\n"
        "Buoc 2: Chay T=20 forward passes cho moi sample\n"
        "  Moi lan dropout tat neurons KHAC NHAU -> ket qua KHAC NHAU\n"
        "  -> Thu duoc T=20 predictions cho moi sample\n\n"
        "Buoc 3: Tinh mean probability\n"
        "  p_mean = (1/T) * sum(p_t)  voi t = 1..T\n\n"
        "Buoc 4: Tinh Predictive Entropy (do bat dinh)\n"
        "  H = -[p_mean * log(p_mean) + (1-p_mean) * log(1-p_mean)]\n\n"
        "  H thap (gan 0): Mo hinh chac chan -> TIN CAY\n"
        "  H cao (gan 0.693 = ln(2)): Mo hinh khong chac chan -> KHONG TIN CAY"
    )

    doc.add_heading("7.3. UncertaintyFilteredStrategy", level=2)
    doc.add_paragraph(
        "Dieu kien ACCEPT pseudo-label (phai thoa CA HAI):\n"
        "  1. p_mean[i,c] > confidence_threshold (0.7)\n"
        "  2. H[i,c] < uncertainty_threshold (0.5)\n\n"
        "So sanh voi PositiveOnlyStrategy:\n"
        "  PositiveOnly: chi can prob > 0.7 -> ACCEPT\n"
        "  UncertaintyFiltered: can prob > 0.7 VA entropy < 0.5 -> ACCEPT\n\n"
        "Ket qua: UncertaintyFiltered se REJECT nhieu hon,\n"
        "nhung nhung pseudo-label con lai se CHINH XAC hon\n"
        "-> Quality over quantity"
    )

    doc.add_heading("7.4. Chi phi tinh toan", level=2)
    doc.add_paragraph(
        "MC Dropout can T=20 forward passes thay vi 1:\n"
        "  - Teacher inference: ~7 phut (1 pass)\n"
        "  - MC Dropout: ~7 * 20 = ~140 phut (20 passes)\n\n"
        "Day la trade-off: ton nhieu thoi gian hon nhung pseudo-labels tot hon."
    )


def build_section8(doc):
    doc.add_heading("8. EVALUATION - DANH GIA MO HINH", level=1)

    doc.add_heading("8.1. AUROC la gi?", level=2)
    doc.add_paragraph(
        "AUROC = Area Under the Receiver Operating Characteristic Curve\n\n"
        "ROC Curve:\n"
        "  - Truc X: False Positive Rate (FPR) = FP / (FP + TN)\n"
        "  - Truc Y: True Positive Rate (TPR) = TP / (TP + FN)\n"
        "  - Ve duong cong khi thay doi threshold tu 0 den 1\n\n"
        "AUROC:\n"
        "  - Dien tich duoi duong ROC\n"
        "  - Gia tri tu 0 den 1 (1 = hoan hao, 0.5 = random)\n"
        "  - KHONG phu thuoc vao threshold -> phu hop cho class imbalanced\n\n"
        "Tai sao dung AUROC thay vi Accuracy?\n"
        "  Vi du lieu CheXpert rat imbalanced:\n"
        "  Pneumothorax chi co ~5% positive -> mo hinh du doan tat ca = 0\n"
        "  van dat accuracy 95% nhung AUROC chi 0.5 (vo dung)"
    )

    doc.add_heading("8.2. Macro AUROC vs Per-class AUROC", level=2)
    doc.add_paragraph(
        "Per-class AUROC: Tinh AUROC rieng cho tung class\n"
        "  VD: Cardiomegaly AUROC = 0.77, Pleural Effusion = 0.92\n\n"
        "Macro AUROC: Trung binh cong cua tat ca per-class AUROC\n"
        "  macro_auroc = mean([0.77, 0.92, 0.69, 0.88, 0.77, 0.89]) = 0.82\n\n"
        "Ket qua Supervised Baseline (epoch 6):\n"
        "  Macro AUROC: 0.8195\n"
        "  Cardiomegaly:      0.7708\n"
        "  Pleural Effusion:  0.9224\n"
        "  Pneumothorax:      0.6894  <- Yeu nhat\n"
        "  Consolidation:     0.8754\n"
        "  Atelectasis:       0.7670\n"
        "  Edema:             0.8921"
    )

    doc.add_heading("8.3. Optional Metrics", level=2)
    doc.add_paragraph(
        "Ngoai AUROC, du an ho tro them:\n"
        "  - Macro F1: F1-score trung binh (can threshold, mac dinh 0.5)\n"
        "  - MAP (Mean Average Precision): trung binh AP cua tat ca class\n\n"
        "Nhung AUROC la metric CHINH vi:\n"
        "  1. Threshold-independent\n"
        "  2. La metric chuan cua CheXpert benchmark\n"
        "  3. Xu ly tot class imbalance"
    )



def build_section9(doc):
    doc.add_heading("9. CAU TRUC CODE CHI TIET", level=1)

    doc.add_heading("9.1. Thu muc du an", level=2)
    doc.add_paragraph(
        "semisub-cxr/\n"
        "  configs/              <- YAML config files cho moi setting\n"
        "    supervised_baseline.yaml\n"
        "    pseudo_label.yaml\n"
        "    uncertainty_filter.yaml\n"
        "    lvmmed_finetune.yaml\n"
        "  scripts/              <- CLI entry points\n"
        "    train.py            <- Huan luyen mo hinh\n"
        "    evaluate.py         <- Danh gia checkpoint\n"
        "    generate_splits.py  <- Tao labeled/unlabeled split\n"
        "    generate_pseudo_labels.py <- Tao pseudo-labels\n"
        "    demo_run.py         <- Demo voi synthetic data\n"
        "  src/                  <- Source code chinh\n"
        "    config.py           <- Dataclass configs + YAML loader\n"
        "    data/\n"
        "      chexpert_dataset.py  <- PyTorch Dataset cho CheXpert\n"
        "      combined_dataset.py  <- Gop labeled + pseudo-labeled\n"
        "      split_generator.py   <- Stratified split generation\n"
        "      transforms.py        <- Image augmentation pipelines\n"
        "    models/\n"
        "      backbone_factory.py  <- Registry pattern cho backbones\n"
        "      classifier.py        <- MultiLabelClassifier\n"
        "    training/\n"
        "      trainer.py           <- Training loop + checkpointing\n"
        "      losses.py            <- Masked BCE loss\n"
        "    pseudo_labeling/\n"
        "      teacher_inference.py    <- Teacher forward pass\n"
        "      mc_dropout.py           <- MC Dropout uncertainty\n"
        "      threshold_strategy.py   <- Accept/reject strategies\n"
        "      artifact_io.py          <- Save/load pseudo-label CSV\n"
        "    evaluation/\n"
        "      metrics.py           <- AUROC, F1, MAP computation\n"
        "      evaluator.py         <- Checkpoint evaluation\n"
        "      reporting.py         <- Charts, tables generation\n"
        "    utils/\n"
        "      seed.py              <- Reproducibility seeds\n"
        "      checkpoint.py        <- Save/load checkpoints\n"
        "      sanity_checks.py     <- Pipeline validation\n"
        "      logging_utils.py     <- Logging + git tracking\n"
        "  outputs/               <- Ket qua (generated)\n"
        "    splits/              <- Split JSON files\n"
        "    checkpoints/         <- Model checkpoints\n"
        "    pseudo_labels/       <- Pseudo-label CSV artifacts\n"
        "    results/             <- Evaluation CSV/JSON\n"
        "    plots/               <- Charts PNG"
    )

    doc.add_heading("9.2. Config System", level=2)
    doc.add_paragraph(
        "Toan bo cau hinh duoc dinh nghia bang Python dataclasses trong src/config.py:\n\n"
        "ExperimentConfig (top-level):\n"
        "  experiment_name: ten thi nghiem\n"
        "  setting: 'supervised' | 'pseudo_label' | 'uncertainty_filter'\n"
        "  data: DataConfig (dataset_path, label_set, uncertain_policy, ...)\n"
        "  split: SplitConfig (labeled_ratio, seed, splits_dir)\n"
        "  model: ModelConfig (backbone, dropout_rate, num_classes)\n"
        "  training: TrainingConfig (lr, batch_size, epochs, optimizer)\n"
        "  pseudo_label: PseudoLabelConfig (enabled, thresholds, policy)\n"
        "  uncertainty: UncertaintyConfig (enabled, mc_passes, metric)\n"
        "  output_dir: thu muc output\n\n"
        "YAML config duoc load bang dacite (strict mode):\n"
        "  config = load_config('configs/supervised_baseline.yaml')\n"
        "  -> Tra ve ExperimentConfig instance voi type checking"
    )

    doc.add_heading("9.3. Checkpoint System", level=2)
    doc.add_paragraph(
        "Moi checkpoint (.pt file) chua:\n"
        "  model_state_dict: weights cua model\n"
        "  optimizer_state_dict: trang thai optimizer (de resume)\n"
        "  epoch: epoch hien tai\n"
        "  config: toan bo experiment config\n"
        "  val_metrics: {macro_auroc, per_class_auroc}\n"
        "  seed: random seed\n"
        "  git_commit: git hash (tracking code version)\n"
        "  library_versions: {torch, numpy, python}\n\n"
        "2 loai checkpoint:\n"
        "  best_checkpoint.pt: model tot nhat (macro_auroc cao nhat)\n"
        "  last_checkpoint.pt: model cuoi cung (de resume training)"
    )

    doc.add_heading("9.4. Sanity Checks", level=2)
    doc.add_paragraph(
        "Du an co 4 sanity checks chay truoc moi buoc:\n\n"
        "1. check_data_loading(dataset):\n"
        "   - Load 5 samples, kiem tra format (image, label, index)\n"
        "   - Kiem tra label shape = (6,)\n\n"
        "2. check_split_validity(split_meta, dataset_size, ratio):\n"
        "   - Labeled va unlabeled KHONG OVERLAP\n"
        "   - Union = toan bo dataset\n"
        "   - Moi class co it nhat 1 positive\n\n"
        "3. check_training_step(model, batch, device):\n"
        "   - 1 forward pass -> loss phai FINITE\n"
        "   - 1 backward pass -> gradient phai NON-ZERO\n\n"
        "4. check_evaluation(evaluator, model, val_loader):\n"
        "   - Chay inference tren 50 samples\n"
        "   - Kiem tra AUROC tinh duoc (khong loi)"
    )



def build_section10(doc):
    doc.add_heading("10. CAU HOI THUYET TRINH THUONG GAP", level=1)

    doc.add_heading("Q1: Tai sao chon semi-supervised thay vi supervised?", level=2)
    doc.add_paragraph(
        "Trong y te, gan nhan du lieu can bac si chuyen mon -> rat dat va cham.\n"
        "Benh vien co hang trieu anh X-quang nhung chi mot phan nho co nhan.\n"
        "Semi-supervised tan dung du lieu khong nhan de cai thien mo hinh\n"
        "ma khong can them nhan thu cong."
    )

    doc.add_heading("Q2: Tai sao dung multi-label thay vi multi-class?", level=2)
    doc.add_paragraph(
        "Mot benh nhan co the co NHIEU benh cung luc:\n"
        "  VD: Cardiomegaly + Pleural Effusion + Edema\n"
        "Multi-class chi cho phep 1 nhan -> khong phu hop.\n"
        "Multi-label: moi class la 1 bai toan binary doc lap."
    )

    doc.add_heading("Q3: Tai sao dung ResNet-50 ma khong dung model khac?", level=2)
    doc.add_paragraph(
        "1. ResNet-50 la backbone chuan trong medical imaging research\n"
        "2. Co san ImageNet pretrained weights (transfer learning)\n"
        "3. LVM-Med cung dua tren ResNet-50 -> de so sanh\n"
        "4. Can bang giua performance va computational cost\n"
        "5. 2048-dim features du manh cho 6-class classification"
    )

    doc.add_heading("Q4: Pseudo-labeling co van de gi?", level=2)
    doc.add_paragraph(
        "1. Confirmation bias: Teacher sai -> Student hoc theo cai sai\n"
        "   -> Giai phap: Threshold cao (0.7) de chi accept nhan tin cay\n\n"
        "2. Class imbalance trong pseudo-labels:\n"
        "   Pleural Effusion: 41K accepted vs Consolidation: 21 accepted\n"
        "   -> Giai phap: Per-class threshold (co the dieu chinh rieng)\n\n"
        "3. Overconfident predictions:\n"
        "   Model du doan prob cao nhung sai\n"
        "   -> Giai phap: MC Dropout uncertainty filtering"
    )

    doc.add_heading("Q5: MC Dropout khac gi voi Dropout binh thuong?", level=2)
    doc.add_paragraph(
        "Dropout binh thuong:\n"
        "  - Training: BAT (regularization)\n"
        "  - Inference: TAT (dung tat ca neurons)\n\n"
        "MC Dropout:\n"
        "  - Training: BAT (nhu binh thuong)\n"
        "  - Inference: VAN BAT (de tao nhieu predictions khac nhau)\n"
        "  - Chay T=20 lan -> tinh mean va variance\n"
        "  - Variance cao = uncertainty cao = khong tin cay\n\n"
        "Co so ly thuyet: Gal & Ghahramani (2016) chung minh rang\n"
        "MC Dropout xap xi Bayesian inference -> uoc luong uncertainty\n"
        "ma khong can thay doi kien truc model."
    )

    doc.add_heading("Q6: Tai sao dung Predictive Entropy thay vi Variance?", level=2)
    doc.add_paragraph(
        "Predictive Entropy: H = -[p*log(p) + (1-p)*log(1-p)]\n"
        "  - Gia tri tu 0 (chac chan) den ln(2)=0.693 (hoan toan khong chac)\n"
        "  - Phu hop cho binary classification (moi class la binary)\n"
        "  - Tinh tren MEAN probability -> phan anh tong uncertainty\n\n"
        "Variance cung la 1 lua chon nhung entropy pho bien hon\n"
        "trong literature va co ly thuyet information theory ho tro."
    )

    doc.add_heading("Q7: Ket qua du kien nhu the nao?", level=2)
    doc.add_paragraph(
        "Du kien (tu thap den cao):\n"
        "  1. Supervised Baseline: ~0.82 macro AUROC (chi 5% data)\n"
        "  2. LVM-Med: ~0.85 (backbone tot hon)\n"
        "  3. Pseudo-Label: ~0.86 (them du lieu tu pseudo-labels)\n"
        "  4. Uncertainty Filter: ~0.87 (pseudo-labels chat luong hon)\n\n"
        "Ket qua thuc te supervised baseline: 0.8195 macro AUROC\n"
        "  -> Kha tot cho chi 5% labeled data!"
    )

    doc.add_heading("Q8: Reproducibility duoc dam bao nhu the nao?", level=2)
    doc.add_paragraph(
        "1. set_all_seeds(42): Set seed cho Python, NumPy, PyTorch, CUDA\n"
        "2. torch.backends.cudnn.deterministic = True\n"
        "3. Split duoc luu JSON -> cung seed = cung split\n"
        "4. Checkpoint luu git_commit + library_versions\n"
        "5. Config duoc luu JSON ben canh checkpoint\n\n"
        "-> Bat ky ai cung co the reproduce ket qua voi cung config + seed."
    )

    doc.add_heading("Q9: Giai thich data flow tu dau den cuoi?", level=2)
    doc.add_paragraph(
        "1. YAML config -> load_config() -> ExperimentConfig dataclass\n"
        "2. CheXpertDataset: CSV + images -> (image_tensor, label_vector, index)\n"
        "3. generate_split(): dataset -> labeled_indices + unlabeled_indices\n"
        "4. Subset(dataset, labeled_indices) -> train_loader\n"
        "5. Trainer.train(): train_loader + val_loader -> best_checkpoint.pt\n"
        "6. run_teacher_inference(): model + unlabeled_loader -> probabilities\n"
        "7. PositiveOnlyStrategy.accept(): probs -> pseudo_labels + rejection_mask\n"
        "8. save_pseudo_label_artifact(): -> pseudo_labels.csv\n"
        "9. CombinedDataset(labeled + pseudo) -> combined_train_loader\n"
        "10. Trainer.train(): combined_loader -> student_checkpoint.pt\n"
        "11. evaluate_checkpoint(): checkpoint -> AUROC metrics\n"
        "12. reporting: metrics -> charts + tables"
    )

    doc.add_heading("Q10: Cac ky thuat Machine Learning su dung?", level=2)
    doc.add_paragraph(
        "1. Transfer Learning: Dung ImageNet/LVM-Med pretrained weights\n"
        "2. Semi-Supervised Learning: Pseudo-labeling (teacher-student)\n"
        "3. Multi-Label Classification: BCE loss, per-class AUROC\n"
        "4. Data Augmentation: RandomFlip, RandomRotation\n"
        "5. Regularization: Dropout (p=0.5), Weight Decay (L2)\n"
        "6. Learning Rate Scheduling: Cosine Annealing\n"
        "7. Bayesian Approximation: MC Dropout for uncertainty\n"
        "8. Stratified Sampling: MultilabelStratifiedShuffleSplit\n"
        "9. Masked Loss: Selective gradient computation\n"
        "10. Checkpointing: Best/last model persistence"
    )




if __name__ == "__main__":
    main()
